from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import copy

BLUE=RGBColor(0x0B,0x2A,0x6F); GOLD=RGBColor(0xB0,0x8D,0x3C)
IMG="decor_unacita_cadre_A4.png"

def add_frame(doc):
    sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(6.6); sec.bottom_margin=Cm(3.8)
    sec.left_margin=Cm(3); sec.right_margin=Cm(3)
    sec.header_distance=Cm(0.5)
    p=sec.header.paragraphs[0]
    run=p.add_run()
    run.add_picture(IMG,width=Cm(21),height=Cm(29.7))
    inline=run._r.xpath('.//wp:inline')[0]
    ns=inline.nsmap
    graphic=inline.find('{http://schemas.openxmlformats.org/drawingml/2006/main}graphic')
    docPr=inline.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr')
    anchor=parse_xml(f'''<wp:anchor {nsdecls("wp","a","pic","r")} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
<wp:simplePos x="0" y="0"/>
<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
<wp:extent cx="{Cm(21)}" cy="{Cm(29.7)}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>
<wp:docPr id="{docPr.get('id')}" name="Cadre UNACITA"/><wp:cNvGraphicFramePr/></wp:anchor>''')
    anchor.append(copy.deepcopy(graphic))
    inline.getparent().replace(inline,anchor)

def style(doc):
    st=doc.styles['Normal']; st.font.name='Georgia'; st.font.size=Pt(11)
    st.paragraph_format.space_after=Pt(6)

def para(doc,text="",bold=False,size=None,color=None,align=None,italic=False,after=None):
    p=doc.add_paragraph(); r=p.add_run(text); r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=color
    if align is not None: p.alignment=align
    if after is not None: p.paragraph_format.space_after=Pt(after)
    return p

def heading(doc,text):
    p=para(doc,text.upper(),bold=True,size=11.5,color=BLUE,after=3)
    p.paragraph_format.space_before=Pt(10)
    pPr=p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="B08D3C"/></w:pBdr>'))
    return p

def entete(doc):
    C=WD_ALIGN_PARAGRAPH.CENTER
    para(doc,"UNACITA — SECTION DE SAINT-AVÉ",bold=True,size=15,color=BLUE,align=C,after=0)
    para(doc,"Union Nationale des Anciens Combattants d'Indochine, des Théâtres d'opérations extérieurs et d'Afrique du Nord",size=8.5,italic=True,align=C,after=0)
    para(doc,"[Adresse du siège] — 56890 Saint-Avé",size=9,align=C,after=10)

def modele():
    d=Document(); style(d); add_frame(d); entete(d)
    para(d,"Objet : ……………………………………………………",bold=True)
    para(d,"Texte du document …")
    d.save("modele_unacita_saint-ave.docx")

def pv():
    d=Document(); style(d); add_frame(d); entete(d)
    C=WD_ALIGN_PARAGRAPH.CENTER
    para(d,"PROCÈS-VERBAL",bold=True,size=18,color=BLUE,align=C,after=0)
    para(d,"de la réunion du lundi 28 septembre 2026",bold=True,size=12,align=C,after=8)
    para(d,"L'an deux mil vingt-six, le lundi 28 septembre, à [heure], les membres de la section de Saint-Avé de l'UNACITA se sont réunis à [lieu], sur convocation du Président, [Nom Prénom].")
    heading(d,"Présents")
    para(d,"[Nom Prénom, fonction] ; [Nom Prénom, fonction] ; [Nom Prénom, fonction] ; …")
    heading(d,"Excusés / absents")
    para(d,"[Nom Prénom] ; [Nom Prénom] ; …")
    para(d,"Le quorum est [atteint / non atteint] : [X] membres présents ou représentés sur [Y] inscrits.",italic=True)
    heading(d,"Ordre du jour")
    for i,t in enumerate(["Approbation du précédent procès-verbal","Point sur la situation financière","Préparation des cérémonies commémoratives (11 novembre 2026)","Adhésions et vie de la section","Questions diverses"],1):
        para(d,f"{i}. {t}",after=2)
    heading(d,"Délibérations")
    for i,(t,x) in enumerate([("Approbation du précédent procès-verbal","Le procès-verbal de la réunion du [date] est soumis au vote. Adopté à [l'unanimité / X voix pour, Y contre, Z abstentions]."),
        ("Situation financière","Le Trésorier présente [solde, recettes, dépenses]. Après discussion, le bilan est [approuvé]."),
        ("Cérémonies commémoratives","Le Président expose l'organisation de la cérémonie du 11 novembre 2026 : [programme, porte-drapeaux, dépôt de gerbe, coordination avec la mairie]."),
        ("Adhésions et vie de la section","[Nouvelles adhésions, cotisations, démarches en cours]."),
        ("Questions diverses","[Sujets abordés]")],1):
        para(d,f"{i}. {t}",bold=True,after=1)
        para(d,x)
    heading(d,"Clôture")
    para(d,"L'ordre du jour étant épuisé, la séance est levée à [heure]. De tout ce qui précède, il a été dressé le présent procès-verbal, signé par le Président et le Secrétaire.")
    para(d,"Fait à Saint-Avé, le 28 septembre 2026.",after=14)
    t=d.add_table(rows=2,cols=2)
    t.rows[0].cells[0].text="Le Président"; t.rows[0].cells[1].text="Le Secrétaire"
    t.rows[1].cells[0].text="\n\n[Nom Prénom]"; t.rows[1].cells[1].text="\n\n[Nom Prénom]"
    for c in t.rows[0].cells:
        for r in c.paragraphs[0].runs: r.bold=True
    d.save("proces_verbal_2026-09-28.docx")
modele(); pv()
